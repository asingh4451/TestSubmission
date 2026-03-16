from pptx import Presentation
from pptx.util import Inches

def create_sample_input(output_path):
    prs = Presentation()
    
    # Slide 1: Title
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Q3 Product Roadmap"
    subtitle.text = "Strategic Vision & Key Deliverables"

    # Slide 2: Strategy (Bullet points)
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Key Strategic Pillars"
    tf = body_shape.text_frame
    tf.text = "1. Customer Obsession"
    p = tf.add_paragraph()
    p.text = "Focus on retention and UX"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "2. Scalability"
    p = tf.add_paragraph()
    p.text = "Infrastructure upgrades for 10x growth"
    p.level = 1

    # Slide 3: Financials (Table)
    table_slide_layout = prs.slide_layouts[5] # Title only
    slide = prs.slides.add_slide(table_slide_layout)
    slide.shapes.title.text = "Projected Growth"
    
    rows, cols = 3, 2
    left, top, width, height = Inches(2), Inches(2), Inches(6), Inches(1.5)
    table = slide.shapes.add_table(rows, cols, left, top, width, height).table
    
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Target"
    table.cell(1, 0).text = "Revenue"
    table.cell(1, 1).text = "$5M"
    table.cell(2, 0).text = "Users"
    table.cell(2, 1).text = "100K"

    prs.save(output_path)
    print(f"Sample input created at {output_path}")

if __name__ == "__main__":
    create_sample_input("input/test_input.pptx")

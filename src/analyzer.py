import os
from pptx import Presentation

def analyze_template(template_path):
    """
    Analyzes a PPTX template to identify layouts and their placeholder types.
    """
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template not found: {template_path}")

    prs = Presentation(template_path)
    analysis = {
        "template_name": os.path.basename(template_path),
        "layouts": []
    }

    for layout in prs.slide_layouts:
        layout_info = {
            "name": layout.name,
            "placeholders": []
        }
        for ph in layout.placeholders:
            layout_info["placeholders"].append({
                "idx": ph.placeholder_format.idx,
                "type": str(ph.placeholder_format.type),
                "name": ph.name
            })
        analysis["layouts"].append(layout_info)

    return analysis

if __name__ == "__main__":
    import sys
    import json
    if len(sys.argv) < 2:
        print("Usage: python analyzer.py <path_to_template>")
        sys.exit(1)
    
    template_file = sys.argv[1]
    analysis_data = analyze_template(template_file)
    print(json.dumps(analysis_data, indent=4))

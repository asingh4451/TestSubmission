import json
import os
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

def extract_content(pptx_path):
    """
    Extracts content from a PPTX file and returns it as a structured dictionary.
    """
    if not os.path.exists(pptx_path):
        raise FileNotFoundError(f"File not found: {pptx_path}")

    prs = Presentation(pptx_path)
    content = {
        "metadata": {
            "title": os.path.basename(pptx_path),
            "slide_count": len(prs.slides)
        },
        "slides": []
    }

    for i, slide in enumerate(prs.slides):
        slide_data = {
            "slide_index": i + 1,
            "layout_name": slide.slide_layout.name,
            "elements": []
        }

        # Extract title if it exists and is meaningful
        if slide.shapes.title and slide.shapes.title.text.strip():
            slide_data["title"] = slide.shapes.title.text
        else:
            slide_data["title"] = None

        for shape in slide.shapes:
            # Skip known internal/non-visual shapes
            if "think-cell" in shape.name.lower():
                continue
            element = {
                "name": shape.name,
                "type": str(shape.shape_type),
                "left": shape.left.pt if hasattr(shape, "left") else None,
                "top": shape.top.pt if hasattr(shape, "top") else None,
                "width": shape.width.pt if hasattr(shape, "width") else None,
                "height": shape.height.pt if hasattr(shape, "height") else None,
            }

            if shape.has_text_frame:
                element["text_content"] = []
                for paragraph in shape.text_frame.paragraphs:
                    para_data = {
                        "text": paragraph.text,
                        "level": paragraph.level,
                        "style": {
                            "font_name": paragraph.font.name if hasattr(paragraph.font, "name") else None,
                            "font_size": paragraph.font.size.pt if paragraph.font.size else None,
                            "bold": paragraph.font.bold,
                            "italic": paragraph.font.italic
                        }
                    }
                    element["text_content"].append(para_data)

            if shape.has_table:
                element["table_content"] = []
                for row in shape.table.rows:
                    row_data = [cell.text for cell in row.cells]
                    element["table_content"].append(row_data)

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                element["is_image"] = True
                # Images are harder to "extract" to JSON, but we note their presence

            slide_data["elements"].append(element)

        content["slides"].append(slide_data)

    return content

if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("Usage: python extractor.py <path_to_pptx>")
        sys.exit(1)
    
    input_file = sys.argv[1]
    extracted_data = extract_content(input_file)
    output_file = input_file.replace(".pptx", "_content.json")
    
    with open(output_file, 'w') as f:
        json.dump(extracted_data, f, indent=4)
    
    print(f"Content extracted to {output_file}")
